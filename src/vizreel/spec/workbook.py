"""Reading a chart's data from a sheet of an Excel workbook."""

import zipfile
from datetime import date, datetime, time, timedelta
from pathlib import Path

from vizreel.format.locales import Locale
from vizreel.spec.data import Table, TableError

WORKBOOK_SUFFIXES = (".xlsx", ".xlsm")
"""The Excel workbooks a chart can read its data from."""


def read_workbook(path: Path, locale: Locale, sheet: str | None = None) -> Table:
    """Read a sheet of an Excel workbook whose first row names its columns.

    A cell is read as the workbook stores it, not as Excel shows it: a number as its value,
    whatever its format (25% is 0.25), a date as ISO text such as "2024-03-31", and a formula
    as the value Excel last calculated. Empty rows are skipped.

    Args:
        path: The workbook to read.
        locale: How the sheet writes numbers that it stores as text.
        sheet: The name of the sheet to read; the first sheet if None.

    Raises:
        TableError: The workbook cannot be read, lacks the sheet, has no rows, or has a cell
            right of the last column the header names.
    """
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except FileNotFoundError:
        raise TableError(f"was not found in {path.parent}") from None
    except PermissionError:
        raise TableError("cannot be read: permission denied") from None
    except (zipfile.BadZipFile, KeyError, OSError):
        raise TableError("is not an Excel workbook; save it again as .xlsx") from None
    try:
        if sheet is not None and sheet not in workbook.sheetnames:
            names = ", ".join(f'"{name}"' for name in workbook.sheetnames)
            raise TableError(f'has no sheet "{sheet}". Sheets: {names}')
        worksheet = workbook[sheet] if sheet is not None else workbook.worksheets[0]
        rows = [
            (number, [cell_text(value) for value in values])
            for number, values in enumerate(worksheet.iter_rows(min_row=1, values_only=True), 1)
        ]
    finally:
        workbook.close()
    return _table(
        [(number, cells) for number, cells in rows if any(text for text, _ in cells)], locale
    )


def cell_text(value: object) -> tuple[str, float | None]:
    """A cell's text, as a table holds it, and its number if the workbook stores one.

    A whole number is written without decimals, so that a year stored as a number reads
    "2024"; a date is written in ISO form, with the time of day unless it is midnight.
    """
    if value is None:
        return "", None
    if isinstance(value, bool):
        return ("TRUE" if value else "FALSE"), None
    if isinstance(value, int | float):
        number = float(value)
        return (str(int(number)) if number.is_integer() else repr(number)), number
    if isinstance(value, datetime):
        if value.time() == time():
            return value.date().isoformat(), None
        return value.isoformat(sep=" ", timespec="minutes"), None
    if isinstance(value, date | time):
        return value.isoformat(), None
    if isinstance(value, timedelta):
        return str(value), None
    return str(value).strip(), None


def _table(rows: list[tuple[int, list[tuple[str, float | None]]]], locale: Locale) -> Table:
    from openpyxl.utils import get_column_letter

    if not rows:
        raise TableError("is empty; the first row names the columns")
    (_, header_cells), body = rows[0], rows[1:]
    header = [text for text, _ in header_cells]
    while header and not header[-1]:
        header.pop()
    if not body:
        raise TableError("has a header row but no rows of data")
    texts, numbers, lines = [], [], []
    for number, cells in body:
        extra = next((index for index in range(len(header), len(cells)) if cells[index][0]), None)
        if extra is not None:
            raise TableError(
                "has a cell right of the last column the header names, in column "
                f"{get_column_letter(extra + 1)}; name its column in the first row or empty it",
                row=number,
            )
        cells = (cells + [("", None)] * len(header))[: len(header)]
        texts.append([text for text, _ in cells])
        numbers.append([value for _, value in cells])
        lines.append(number)
    return Table(header, texts, lines, locale, numbers)
